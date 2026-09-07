import io
import json
import zipfile
from pathlib import Path

import pytest

from app.providers import MockProvider
from app.skills import (
    MULTI_FORMAT_OUTPUT, SkillPackageError, SkillPackageStore, SkillQuestion, SkillRunService,
    _extract_json, _fallback_spec, _parse_skill_md, show_if_met,
)

REPO_ROOT = Path(__file__).resolve().parent.parent
SKILLS_DIR = REPO_ROOT / "skills"


def _store(tmp_path: Path) -> SkillPackageStore:
    # Every test gets its own tmp_path-backed data_dir (never the bare
    # no-arg default) so test runs never write into the real repo's
    # data/skills/ — same convention tests/test_storage.py uses throughout
    # for SessionStore.
    return SkillPackageStore(data_dir=tmp_path / "skills")


def _runs(store: SkillPackageStore, tmp_path: Path, provider=None) -> SkillRunService:
    return SkillRunService(store, provider or MockProvider(), data_dir=tmp_path / "skill-runs")


# --- built-in skill loading ---------------------------------------------------

def test_store_loads_both_builtin_skills(tmp_path):
    store = _store(tmp_path)
    ids = set(store.skills)
    assert {"ppt-generator", "docx-generator"} <= ids
    ppt = store.get("ppt-generator")
    assert ppt.builtin is True
    assert ppt.output == "pptx"
    assert any(q.id == "topic" and q.required for q in ppt.questions)
    docx = store.get("docx-generator")
    assert docx.output == "docx"
    assert any(q.id == "doc_type" for q in docx.questions)


def test_list_and_get_and_missing_skill(tmp_path):
    store = _store(tmp_path)
    listed = store.list()
    assert any(s["skill_id"] == "ppt-generator" for s in listed)
    with pytest.raises(KeyError):
        store.get("does-not-exist")


# --- SKILL.md parsing ---------------------------------------------------------

def _write_skill_md(tmp_path: Path, body: str) -> Path:
    skill_dir = tmp_path / "a-skill"
    skill_dir.mkdir()
    path = skill_dir / "SKILL.md"
    path.write_text(body, encoding="utf-8")
    return path


def test_parse_skill_md_requires_frontmatter(tmp_path):
    path = _write_skill_md(tmp_path, "no frontmatter here")
    with pytest.raises(SkillPackageError, match="front matter"):
        _parse_skill_md(path, skill_id="a-skill", builtin=False)


def test_parse_skill_md_requires_name(tmp_path):
    path = _write_skill_md(tmp_path, "---\ndescription: x\n---\nbody")
    with pytest.raises(SkillPackageError, match="name"):
        _parse_skill_md(path, skill_id="a-skill", builtin=False)


def test_parse_skill_md_requires_question_id_and_prompt(tmp_path):
    body = "---\nname: test\nquestions:\n  - prompt: missing id\n---\nbody"
    path = _write_skill_md(tmp_path, body)
    with pytest.raises(SkillPackageError, match="id.*prompt|prompt.*id"):
        _parse_skill_md(path, skill_id="a-skill", builtin=False)


def test_parse_skill_md_minimal_valid():
    body = "---\nname: hello\ndescription: says hello\noutput: pptx\n---\nDo the thing."
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "SKILL.md"
        path.write_text(body, encoding="utf-8")
        skill = _parse_skill_md(path, skill_id="hello", builtin=False)
    assert skill.name == "hello"
    assert skill.output == "pptx"
    assert skill.questions == []
    assert skill.instructions == "Do the thing."


def test_parse_skill_md_file_question_with_accept(tmp_path):
    body = (
        "---\nname: importer\nquestions:\n"
        "  - id: dataset\n    prompt: \"Upload your data\"\n    type: file\n"
        "    accept: [\"csv\", \"xlsx\"]\n    required: true\n---\nbody"
    )
    path = _write_skill_md(tmp_path, body)
    skill = _parse_skill_md(path, skill_id="importer", builtin=False)
    q = skill.questions[0]
    assert q.type == "file"
    assert q.accept == ["csv", "xlsx"]
    assert q.required is True


# --- zip upload / extraction --------------------------------------------------

def _zip_skill_folder(src: Path, top_level: str = "") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for f in src.rglob("*"):
            if f.is_file():
                rel = f.relative_to(src).as_posix()
                arcname = f"{top_level}/{rel}" if top_level else rel
                zf.write(f, arcname=arcname)
    return buf.getvalue()


def test_add_from_zip_round_trip_and_runs(tmp_path):
    store = _store(tmp_path)
    zip_bytes = _zip_skill_folder(SKILLS_DIR / "ppt-generator", top_level="ppt-generator")
    skill = store.add_from_zip(zip_bytes)
    assert skill.builtin is False
    assert skill.name == "ppt-generator"
    assert skill.skill_id in store.skills
    store.delete(skill.skill_id)
    assert skill.skill_id not in store.skills


def test_add_from_zip_rejects_bad_zip(tmp_path):
    store = _store(tmp_path)
    with pytest.raises(SkillPackageError, match="not a valid"):
        store.add_from_zip(b"definitely not a zip file")


def test_add_from_zip_rejects_missing_skill_md(tmp_path):
    store = _store(tmp_path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "no SKILL.md in here")
    with pytest.raises(SkillPackageError, match="SKILL.md"):
        store.add_from_zip(buf.getvalue())


def test_add_from_zip_rejects_zip_slip(tmp_path):
    store = _store(tmp_path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../../evil.txt", "escape attempt")
    with pytest.raises(SkillPackageError, match="Unsafe path"):
        store.add_from_zip(buf.getvalue())


def test_delete_protects_builtin_skills(tmp_path):
    store = _store(tmp_path)
    with pytest.raises(SkillPackageError, match="Built-in"):
        store.delete("ppt-generator")


# --- persistence: uploaded skills survive a restart --------------------------

def test_uploaded_skill_survives_new_store_instance(tmp_path):
    skills_dir = tmp_path / "skills"
    store = SkillPackageStore(data_dir=skills_dir)
    zip_bytes = _zip_skill_folder(SKILLS_DIR / "ppt-generator", top_level="ppt-generator")
    skill = store.add_from_zip(zip_bytes)

    # A fresh instance (simulating a process restart) must reload it from disk.
    reloaded_store = SkillPackageStore(data_dir=skills_dir)
    reloaded = reloaded_store.get(skill.skill_id)
    assert reloaded.name == "ppt-generator"
    assert reloaded.builtin is False


def test_delete_removes_persisted_skill_directory(tmp_path):
    skills_dir = tmp_path / "skills"
    store = SkillPackageStore(data_dir=skills_dir)
    zip_bytes = _zip_skill_folder(SKILLS_DIR / "ppt-generator", top_level="ppt-generator")
    skill = store.add_from_zip(zip_bytes)
    store.delete(skill.skill_id)
    # No leftover directory (empty or otherwise) for the deleted skill.
    assert not (skills_dir / skill.skill_id).exists()


# --- routing ---------------------------------------------------------------

def test_select_for_task_matches_ppt_keyword(tmp_path):
    store = _store(tmp_path)
    skill = store.select_for_task("can you make a powerpoint about our roadmap")
    assert skill is not None and skill.skill_id == "ppt-generator"


def test_select_for_task_returns_none_for_unrelated_text(tmp_path):
    store = _store(tmp_path)
    assert store.select_for_task("what's the weather like today") is None


def test_select_for_chat_matches_precise_phrases(tmp_path):
    store = _store(tmp_path)
    assert store.select_for_chat("can you create a docx for me").skill_id == "docx-generator"
    assert store.select_for_chat("I need a word document about onboarding").skill_id == "docx-generator"
    assert store.select_for_chat("write a proposal for the new tool").skill_id == "docx-generator"
    # Deck requests now route to the richer `pptx` skill — ppt-generator's own
    # chat_triggers are deliberately cleared (see skills/ppt-generator/SKILL.md)
    # so this is deterministic, not first-match-wins over dict order.
    assert store.select_for_chat("make me a powerpoint about Q3 results").skill_id == "pptx"
    assert store.select_for_chat("can you build a slide deck").skill_id == "pptx"


def test_select_for_chat_does_not_false_positive_on_common_words(tmp_path):
    # "word" and "report" alone are common enough that a bare keyword match
    # (select_for_task's approach) would misfire here -- select_for_chat must not.
    store = _store(tmp_path)
    assert store.select_for_chat("what's a good word for happy?") is None
    assert store.select_for_chat("can you report back on the test results") is None
    assert store.select_for_chat("what's the weather like today") is None


# --- spec drafting / fallback ------------------------------------------------

def test_extract_json_handles_fenced_and_bare_json():
    assert _extract_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert _extract_json('{"a": 1}') == {"a": 1}
    assert _extract_json('here you go: {"a": 1} thanks') == {"a": 1}
    assert _extract_json("not json at all") is None


def test_fallback_spec_builds_usable_ppt_spec(tmp_path):
    store = _store(tmp_path)
    skill = store.get("ppt-generator")
    spec = _fallback_spec(skill, {"topic": "Onboarding", "audience": "New hires", "tone": "Casual"})
    assert spec["title"] == "Onboarding"
    assert spec["slides"] and spec["slides"][0]["bullets"]


def test_fallback_spec_builds_usable_docx_spec(tmp_path):
    store = _store(tmp_path)
    skill = store.get("docx-generator")
    spec = _fallback_spec(skill, {"topic": "Policy update", "doc_type": "Memo", "tone": "Formal"})
    assert spec["title"] == "Policy update"
    assert spec["sections"] and spec["sections"][0]["heading"] == "Memo"


# --- SkillRunService: the full ask -> answer -> generate flow ----------------

@pytest.mark.asyncio
async def test_skill_run_full_flow_ppt(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)  # MockProvider's text is never valid JSON
    run = runs.start("ppt-generator")
    assert run.status == "AWAITING_ANSWERS"

    completed = await runs.submit_answers(run.run_id, {
        "topic": "Team offsite recap", "audience": "My team", "tone": "Casual",
    })
    assert completed.status == "COMPLETED"
    assert completed.error is None

    content, filename, content_type = runs.get_output(run.run_id)
    assert content and len(content) > 0
    assert filename.endswith(".pptx")
    assert content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert completed.spec is not None and completed.spec["title"]


@pytest.mark.asyncio
async def test_skill_run_regenerate_from_edited_spec(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")
    completed = await runs.submit_answers(run.run_id, {
        "topic": "Team offsite recap", "audience": "My team", "tone": "Casual",
    })
    assert completed.status == "COMPLETED"
    original_key = completed.output_keys[0]

    edited_spec = dict(completed.spec)
    edited_spec["title"] = "Edited Title From User"
    edited_spec["slides"] = [{"title": "New slide", "bullets": ["Edited bullet one", "Edited bullet two"]}]
    regenerated = runs.regenerate(run.run_id, edited_spec)

    assert regenerated.status == "COMPLETED"
    assert regenerated.spec["title"] == "Edited Title From User"
    assert regenerated.output_keys[0] != original_key  # a fresh file, not a mutation of the old one

    from pptx import Presentation
    content, _, _ = runs.get_output(run.run_id)
    prs = Presentation(io.BytesIO(content))
    assert prs.slides[0].shapes.title.text == "Edited Title From User"
    assert prs.slides[1].shapes.title.text == "New slide"


@pytest.mark.asyncio
async def test_skill_run_regenerate_before_first_generation_rejected(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")  # still AWAITING_ANSWERS
    with pytest.raises(SkillPackageError, match="wait for it to finish"):
        runs.regenerate(run.run_id, {"title": "x", "slides": []})


@pytest.mark.asyncio
async def test_skill_run_full_flow_docx(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("docx-generator")

    completed = await runs.submit_answers(run.run_id, {
        "topic": "New hire policy", "doc_type": "Memo", "audience": "My team", "tone": "Formal",
    })
    assert completed.status == "COMPLETED"
    content, filename, content_type = runs.get_output(run.run_id)
    assert content and len(content) > 0
    assert filename.endswith(".docx")
    assert content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.mark.asyncio
async def test_skill_run_rejects_missing_required_answers(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")
    with pytest.raises(SkillPackageError, match="Missing required answers"):
        await runs.submit_answers(run.run_id, {"topic": "Only the topic"})
    # the run must stay awaiting answers, not fall into a broken GENERATING state
    assert runs.get(run.run_id).status == "AWAITING_ANSWERS"


@pytest.mark.asyncio
async def test_skill_run_download_before_completion_raises(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")
    with pytest.raises(SkillPackageError, match="no completed output"):
        runs.get_output(run.run_id)


def test_skill_run_unknown_run_id_raises_keyerror(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    with pytest.raises(KeyError):
        runs.get("does-not-exist")


# --- SkillRunService: persistence across restarts -----------------------------

@pytest.mark.asyncio
async def test_completed_run_survives_new_service_instance(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")
    completed = await runs.submit_answers(run.run_id, {
        "topic": "Team offsite recap", "audience": "My team", "tone": "Casual",
    })

    # A fresh instance (simulating a process restart) must reload the run
    # record AND still resolve its generated output file.
    reloaded_runs = _runs(store, tmp_path)
    reloaded = reloaded_runs.get(completed.run_id)
    assert reloaded.status == "COMPLETED"
    content, filename, _ = reloaded_runs.get_output(completed.run_id)
    assert content and len(content) > 0
    assert filename.endswith(".pptx")


# --- file-type questions -----------------------------------------------------

def _write_file_skill(tmp_path: Path, accept: list[str]) -> Path:
    """A minimal skill with one required file question, whose generation
    script just proves the resolved path was actually handed to it (writes
    the uploaded file's byte length into the output)."""
    skill_dir = tmp_path / "file-skill"
    scripts_dir = skill_dir / "scripts"
    scripts_dir.mkdir(parents=True)
    accept_yaml = json.dumps(accept)
    (skill_dir / "SKILL.md").write_text(
        "---\nname: file-skill\noutput: txt\nquestions:\n"
        "  - id: dataset\n    prompt: \"Upload a file\"\n    type: file\n"
        f"    accept: {accept_yaml}\n    required: true\n---\nbody",
        encoding="utf-8",
    )
    (scripts_dir / "generate_txt.py").write_text(
        "import json, sys\n"
        "spec = json.loads(sys.stdin.read())\n"
        "out = sys.argv[sys.argv.index('--output') + 1]\n"
        "uploaded = spec.get('uploaded_files', {}).get('dataset')\n"
        "with open(uploaded['path'], 'rb') as f:\n"
        "    size = len(f.read())\n"
        "with open(out, 'w') as f:\n"
        "    f.write(f\"filename={uploaded['filename']} size={size}\")\n",
        encoding="utf-8",
    )
    return skill_dir


def test_upload_answer_file_round_trip_and_generation(tmp_path):
    skill_dir = _write_file_skill(tmp_path, accept=["csv"])
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="file-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    raw = b"a,b,c\n1,2,3\n"
    result = runs.upload_answer_file(run.run_id, "dataset", "data.csv", raw)
    assert result["filename"] == "data.csv"
    assert result["size_bytes"] == len(raw)

    resolved_path = runs.files.resolve(run.run_id, result["file_id"])
    assert resolved_path.read_bytes() == raw


@pytest.mark.asyncio
async def test_file_question_answer_reaches_generation_script(tmp_path):
    skill_dir = _write_file_skill(tmp_path, accept=["csv"])
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="file-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    raw = b"a,b,c\n1,2,3\n"
    uploaded = runs.upload_answer_file(run.run_id, "dataset", "data.csv", raw)

    completed = await runs.submit_answers(run.run_id, {"dataset": uploaded["file_id"]})
    assert completed.status == "COMPLETED", completed.error
    content, _, _ = runs.get_output(run.run_id)
    assert content.decode("utf-8") == f"filename=data.csv size={len(raw)}"


def test_upload_answer_file_rejects_disallowed_extension(tmp_path):
    skill_dir = _write_file_skill(tmp_path, accept=["csv"])
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="file-skill"))
    runs = _runs(store, tmp_path)
    run = runs.start(skill.skill_id)
    with pytest.raises(SkillPackageError, match="allowed types"):
        runs.upload_answer_file(run.run_id, "dataset", "data.exe", b"not allowed")


def test_upload_answer_file_rejects_unknown_question(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")  # has no file-type question
    with pytest.raises(SkillPackageError, match="not a file question"):
        runs.upload_answer_file(run.run_id, "topic", "data.csv", b"x")


@pytest.mark.asyncio
async def test_submit_answers_rejects_expired_file_id(tmp_path):
    skill_dir = _write_file_skill(tmp_path, accept=["csv"])
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="file-skill"))
    runs = _runs(store, tmp_path)
    run = runs.start(skill.skill_id)
    with pytest.raises(SkillPackageError, match="missing or expired"):
        await runs.submit_answers(run.run_id, {"dataset": "not-a-real-file-id"})
    assert runs.get(run.run_id).status == "AWAITING_ANSWERS"


# --- delivery="inline" (see app/skill_render.py, TurnPlan.delivery) -----------

@pytest.mark.asyncio
async def test_submit_answers_inline_delivery_renders_text_no_file(tmp_path):
    # docx-generator has a renderer — delivery="inline" must render the spec
    # as text and skip generation/upload entirely (no output_keys at all).
    # MockProvider's own drafting behavior (used here, same as every other
    # test in this module) isn't the point of this test — just that
    # whatever spec got drafted was rendered as text, not generated as a file.
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("docx-generator")
    completed = await runs.submit_answers(
        run.run_id,
        {"topic": "Q3 performance", "doc_type": "Report", "audience": "Executives", "tone": "Formal"},
        delivery="inline",
    )
    assert completed.status == "COMPLETED_INLINE", completed.error
    assert completed.rendered_text
    assert completed.rendered_text.startswith("# ")  # render_spec_as_chat_text's title heading
    assert completed.output_keys == []
    public = completed.public()
    assert public["download_ready"] is False
    assert public["outputs"] == []
    assert public["rendered_text"] == completed.rendered_text


@pytest.mark.asyncio
async def test_submit_answers_inline_delivery_falls_back_to_file_when_no_renderer(tmp_path):
    # ppt-generator has no renderer yet (see app/skill_render.py) —
    # delivery="inline" must fall back to generating the real file exactly
    # as if delivery had been "file" all along, not silently produce nothing.
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("ppt-generator")
    completed = await runs.submit_answers(
        run.run_id,
        {"topic": "Q3 performance", "audience": "Executives", "tone": "Formal"},
        delivery="inline",
    )
    assert completed.status == "COMPLETED", completed.error
    assert completed.rendered_text is None
    assert completed.output_keys
    assert completed.public()["download_ready"] is True


@pytest.mark.asyncio
async def test_submit_answers_default_delivery_is_file(tmp_path):
    # Every existing caller (including the Skills tab's own route, which
    # passes no delivery argument at all) must keep generating a real file.
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)
    run = runs.start("docx-generator")
    completed = await runs.submit_answers(
        run.run_id,
        {"topic": "Q3 performance", "doc_type": "Report", "audience": "Executives", "tone": "Formal"},
    )
    assert completed.status == "COMPLETED", completed.error
    assert completed.rendered_text is None
    assert completed.output_keys


# --- multi-output support (MULTI_FORMAT_OUTPUT) ------------------------------

def _write_multi_format_skill(tmp_path: Path) -> Path:
    """A minimal skill.output == MULTI_FORMAT_OUTPUT skill whose script
    writes "<output>.<ext>" for each requested spec['output_formats'] entry
    — the exact contract app/skills.py's run_generation_script expects.
    Uses `project_name` as its title question id (rather than a generic
    "title") specifically so MockProvider's fallback path — _fallback_spec's
    MULTI_FORMAT_OUTPUT branch, _fallback_brd_spec — populates something
    real into spec.meta.project_name for this test skill too, without this
    test needing its own bespoke fallback-spec branch."""
    skill_dir = tmp_path / "multi-skill"
    scripts_dir = skill_dir / "scripts"
    scripts_dir.mkdir(parents=True)
    (skill_dir / "SKILL.md").write_text(
        f"---\nname: multi-skill\noutput: {MULTI_FORMAT_OUTPUT}\nquestions:\n"
        "  - id: project_name\n    prompt: \"Title\"\n    type: text\n    required: true\n---\nbody",
        encoding="utf-8",
    )
    (scripts_dir / "generate_multi.py").write_text(
        "import json, sys\n"
        "spec = json.loads(sys.stdin.read())\n"
        "out = sys.argv[sys.argv.index('--output') + 1]\n"
        "title = (spec.get('meta') or {}).get('project_name', '')\n"
        "for fmt in spec.get('output_formats') or ['docx']:\n"
        "    with open(f'{out}.{fmt}', 'w') as f:\n"
        "        f.write(f\"{fmt}:{title}\")\n",
        encoding="utf-8",
    )
    return skill_dir


@pytest.mark.asyncio
async def test_multi_format_skill_produces_both_files(tmp_path):
    skill_dir = _write_multi_format_skill(tmp_path)
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="multi-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    completed = await runs.submit_answers(run.run_id, {"project_name": "Hello", "output_format": "Both"})
    assert completed.status == "COMPLETED", completed.error
    assert len(completed.output_keys) == 2
    assert {k.rsplit(".", 1)[-1] for k in completed.output_keys} == {"docx", "pdf"}

    public = completed.public()
    assert set(public["outputs"]) == {"docx", "pdf"}

    docx_content, docx_name, _ = runs.get_output(run.run_id, format="docx")
    pdf_content, pdf_name, _ = runs.get_output(run.run_id, format="pdf")
    assert docx_content.decode("utf-8") == "docx:Hello"
    assert pdf_content.decode("utf-8") == "pdf:Hello"
    assert docx_name.endswith(".docx") and pdf_name.endswith(".pdf")

    outputs = runs.list_output_keys(run.run_id)
    assert {k.rsplit(".", 1)[-1] for k, _ in outputs} == {"docx", "pdf"}


@pytest.mark.asyncio
async def test_multi_format_skill_docx_only(tmp_path):
    skill_dir = _write_multi_format_skill(tmp_path)
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="multi-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    completed = await runs.submit_answers(run.run_id, {"project_name": "Hello", "output_format": "Word (.docx)"})
    assert completed.status == "COMPLETED"
    assert len(completed.output_keys) == 1
    assert completed.output_keys[0].endswith(".docx")

    with pytest.raises(SkillPackageError, match="no 'pdf' output"):
        runs.get_output(run.run_id, format="pdf")


@pytest.mark.asyncio
async def test_multi_format_skill_regenerate_produces_fresh_files(tmp_path):
    skill_dir = _write_multi_format_skill(tmp_path)
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="multi-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    completed = await runs.submit_answers(run.run_id, {"project_name": "Hello", "output_format": "Both"})
    original_keys = set(completed.output_keys)

    regenerated = runs.regenerate(run.run_id, {**completed.spec, "output_formats": ["docx", "pdf"]})
    assert regenerated.status == "COMPLETED"
    assert set(regenerated.output_keys).isdisjoint(original_keys)  # fresh files, not mutated in place


def test_ordinary_single_output_skill_unaffected(tmp_path):
    # docx-generator/ppt-generator still produce exactly one file, wrapped
    # in a single-element output_keys list — regression check that Part 1's
    # generalization didn't change ordinary skills' behavior.
    store = _store(tmp_path)
    skill = store.get("ppt-generator")
    assert skill.output == "pptx"
    assert skill.output != MULTI_FORMAT_OUTPUT


# --- multiselect + show_if -----------------------------------------------------

def test_show_if_met_true_when_no_condition():
    q = SkillQuestion(id="q1", prompt="x", show_if=None)
    assert show_if_met(q, {}) is True


def test_show_if_met_includes_against_multiselect_answer():
    q = SkillQuestion(id="q2", prompt="x", show_if={"question_id": "sections", "includes": "ROI"})
    assert show_if_met(q, {"sections": json.dumps(["ROI", "SWOT"])}) is True
    assert show_if_met(q, {"sections": json.dumps(["SWOT"])}) is False
    assert show_if_met(q, {}) is False


def test_show_if_met_equals_against_select_answer():
    q = SkillQuestion(id="q3", prompt="x", show_if={"question_id": "output_format", "equals": "PDF"})
    assert show_if_met(q, {"output_format": "pdf"}) is True  # case-insensitive
    assert show_if_met(q, {"output_format": "Word"}) is False


def _write_show_if_skill(tmp_path: Path) -> Path:
    """A skill with a multiselect driver ("sections") and one required
    follow-up question gated by show_if on that driver."""
    skill_dir = tmp_path / "show-if-skill"
    scripts_dir = skill_dir / "scripts"
    scripts_dir.mkdir(parents=True)
    (skill_dir / "SKILL.md").write_text(
        "---\nname: show-if-skill\noutput: txt\nquestions:\n"
        "  - id: sections\n    prompt: \"Sections\"\n    type: multiselect\n"
        "    required: true\n    options: [\"A\", \"B\"]\n"
        "  - id: a_detail\n    prompt: \"A detail\"\n    type: text\n"
        "    required: true\n    show_if: {question_id: sections, includes: \"A\"}\n"
        "---\nbody",
        encoding="utf-8",
    )
    (scripts_dir / "generate_txt.py").write_text(
        "import sys\nout = sys.argv[sys.argv.index('--output') + 1]\n"
        "open(out, 'w').write('ok')\n",
        encoding="utf-8",
    )
    return skill_dir


@pytest.mark.asyncio
async def test_submit_answers_skips_required_check_for_hidden_question(tmp_path):
    skill_dir = _write_show_if_skill(tmp_path)
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="show-if-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    # "A" not selected -> a_detail's show_if is unmet -> its required-ness
    # must be skipped, not block submission even though a_detail is empty.
    completed = await runs.submit_answers(run.run_id, {"sections": json.dumps(["B"])})
    assert completed.status == "COMPLETED", completed.error


@pytest.mark.asyncio
async def test_submit_answers_enforces_required_check_for_shown_question(tmp_path):
    skill_dir = _write_show_if_skill(tmp_path)
    store = _store(tmp_path)
    skill = store.add_from_zip(_zip_skill_folder(skill_dir, top_level="show-if-skill"))
    runs = _runs(store, tmp_path)

    run = runs.start(skill.skill_id)
    # "A" selected -> a_detail's show_if IS met -> its required-ness applies.
    with pytest.raises(SkillPackageError, match="Missing required answers"):
        await runs.submit_answers(run.run_id, {"sections": json.dumps(["A"])})


# --- brd-prd-generator built-in skill ------------------------------------------

def test_brd_prd_generator_loads_as_builtin(tmp_path):
    store = _store(tmp_path)
    skill = store.get("brd-prd-generator")
    assert skill.builtin is True
    assert skill.output == MULTI_FORMAT_OUTPUT
    assert any(q.id == "sections" and q.type == "multiselect" for q in skill.questions)
    roi_question = next(q for q in skill.questions if q.id == "roi_inputs")
    assert roi_question.show_if == {"question_id": "sections", "includes": "ROI / Cost-Benefit Analysis"}


def test_brd_prd_generator_script_tolerates_real_llm_key_variations(tmp_path):
    """Regression test for a real bug found via live manual testing against
    a configured Ollama model: the LLM's JSON matched the documented shape
    loosely, not exactly — human-readable section labels as dict keys
    instead of the short internal ones, cost_items/benefit_items instead
    of costs/benefits, requirement_id instead of req_id, meta.stakeholders
    as a list instead of a string, elicitation_summary as a nested object
    instead of a string. Before the fix (generate_brd.py's _normalize_spec/
    _normalize_roi/_normalize_traceability), this produced a docx with
    ZERO tables — the ROI and traceability sections were silently dropped.
    This spec is the actual (redacted) shape that model returned."""
    import subprocess
    import sys

    script = SKILLS_DIR / "brd-prd-generator" / "scripts" / "generate_brd.py"
    real_world_spec = {
        "meta": {
            "project_name": "Order Portal Revamp",
            "objective": "Speed up order entry for support staff.",
            "stakeholders": ["Support Ops", "IT"],  # list, not a string
            "doc_type": "BRD",
        },
        "elicitation_summary": {  # object, not a string
            "objective": "Speed up order entry for support staff.",
            "stakeholders": ["Support Ops", "IT"],
            "scope": "Revamp the Order Portal for fast search and inventory sync.",
        },
        "sections": {
            "ROI / Cost-Benefit Analysis": {  # human label as the key, not "roi"
                "cost_items": [{"item": "Development cost", "amount": "$40,000", "currency": "USD"}],
                "benefit_items": [{"item": "Savings", "amount": "$15,000", "currency": "USD per year"}],
                "payback_period_years": 2.67,
            },
            "Requirements Traceability Matrix": {  # human label, not "traceability"
                "rows": [
                    {"requirement_id": "REQ-001", "description": "Fast order search",
                     "source": "Business requirement", "test_case_id": "TC-001", "status": "Not Started"},
                ],
            },
        },
        "output_formats": ["docx"],
    }
    output_base = tmp_path / "out"
    proc = subprocess.run(
        [sys.executable, str(script), "--output", str(output_base)],
        input=json.dumps(real_world_spec), capture_output=True, text=True, timeout=30,
    )
    assert proc.returncode == 0, proc.stderr
    docx_path = tmp_path / "out.docx"
    assert docx_path.exists()

    from docx import Document
    doc = Document(docx_path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Support Ops, IT" in all_text  # stakeholders rendered as clean text, not a list repr
    assert "{" not in all_text and "'objective':" not in all_text  # no raw dict repr leaked in

    table_texts = [[c.text for c in row.cells] for t in doc.tables for row in t.rows]
    flat = [" ".join(row) for row in table_texts]
    # ROI costs table + ROI benefits table + traceability table = 3 — all
    # must render despite the key-name mismatch (before the fix: 0 tables).
    assert len(doc.tables) == 3
    assert any("Development cost" in row and "$40,000" in row for row in flat)
    assert any("Savings" in row and "$15,000" in row for row in flat)
    assert any("REQ-001" in row and "Fast order search" in row for row in flat)


@pytest.mark.asyncio
async def test_brd_prd_generator_full_flow_both_formats(tmp_path):
    store = _store(tmp_path)
    runs = _runs(store, tmp_path)  # MockProvider -> exercises _fallback_brd_spec
    run = runs.start("brd-prd-generator")

    completed = await runs.submit_answers(run.run_id, {
        "project_name": "Order Portal Revamp",
        "objective": "Speed up order entry.",
        "stakeholders": "Support Ops",
        "doc_type": "BRD (Business Requirements Document)",
        "output_format": "Both",
        "sections": json.dumps(["ROI / Cost-Benefit Analysis", "Requirements Traceability Matrix"]),
        "roi_inputs": "Dev cost $10k; saves $5k/yr",
        "traceability_scope": "Fast order search",
    })
    assert completed.status == "COMPLETED", completed.error
    assert set(completed.public()["outputs"]) == {"docx", "pdf"}

    docx_content, _, _ = runs.get_output(run.run_id, format="docx")
    pdf_content, _, _ = runs.get_output(run.run_id, format="pdf")
    assert docx_content and len(docx_content) > 0
    assert pdf_content and len(pdf_content) > 0

    from docx import Document
    doc = Document(io.BytesIO(docx_content))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Order Portal Revamp" in all_text
    # ROI + traceability tables actually present
    table_texts = [[c.text for c in row.cells] for t in doc.tables for row in t.rows]
    assert any("Dev cost $10k; saves $5k/yr" in " ".join(row) for row in table_texts)
    assert any("REQ-001" in row for row in table_texts)


# --- pptx built-in skill (richer deck generator) ------------------------------

def test_pptx_skill_loads_correctly(tmp_path):
    store = _store(tmp_path)
    skill = store.get("pptx")
    assert skill.builtin is True
    assert skill.output == "pptx"
    ids = {q.id for q in skill.questions}
    assert {"topic", "audience", "tone", "design", "include_charts", "chart_data", "layout_style"} <= ids
    chart_data_question = next(q for q in skill.questions if q.id == "chart_data")
    assert chart_data_question.show_if == {"question_id": "include_charts", "equals": "Yes"}


def test_select_for_chat_pptx_wins_over_ppt_generator(tmp_path):
    store = _store(tmp_path)
    # ppt-generator's chat_triggers are cleared (see skills/ppt-generator/SKILL.md) —
    # the richer pptx skill is now the deterministic chat-routing target.
    assert store.select_for_chat("make me a powerpoint about Q3").skill_id == "pptx"
    assert store.select_for_chat("can you build a slide deck").skill_id == "pptx"
    ppt = store.get("ppt-generator")
    assert ppt.chat_triggers == []


def _pptx_generate_script() -> Path:
    return SKILLS_DIR / "pptx" / "scripts" / "generate_pptx.py"


def test_generate_pptx_script_produces_real_file(tmp_path):
    import subprocess
    import sys

    from pptx import Presentation

    spec = {
        "title": "Q3 Enterprise Readiness", "subtitle": "Executive Review",
        "tone": "formal", "theme": "midnight_executive",
        "slides": [
            {"layout": "bullets", "title": "Overview", "bullets": ["Revenue up 22%", "Three new logos"], "icon": "check"},
            {"layout": "two_column", "title": "Before vs After",
             "left_heading": "Before", "left_bullets": ["Manual onboarding"],
             "right_heading": "After", "right_bullets": ["Automated onboarding"]},
            {"layout": "stat_callout", "title": "By the numbers",
             "stats": [{"value": "42%", "label": "Faster onboarding"}, {"value": "$1.2M", "label": "New ARR"}]},
            {"layout": "chart", "title": "Quarterly Revenue",
             "chart": {"type": "column", "categories": ["Q1", "Q2", "Q3"], "series": [{"name": "Revenue", "values": [1.0, 1.5, 2.2]}]}},
            {"layout": "section_divider", "title": "Looking Ahead", "subtitle": "Q4 Roadmap"},
        ],
    }
    out_path = tmp_path / "out.pptx"
    proc = subprocess.run(
        [sys.executable, str(_pptx_generate_script()), "--output", str(out_path)],
        input=json.dumps(spec), capture_output=True, text=True, timeout=30,
    )
    assert proc.returncode == 0, proc.stderr
    assert out_path.exists() and out_path.stat().st_size > 0

    prs = Presentation(str(out_path))
    assert len(prs.slides) == 6  # title + 5 content slides — no corruption on reopen


def test_generate_pptx_script_handles_missing_optional_fields(tmp_path):
    import subprocess
    import sys

    from pptx import Presentation

    spec = {"title": "Minimal Deck", "slides": [{"title": "Only slide", "bullets": ["one point"]}]}
    out_path = tmp_path / "minimal.pptx"
    proc = subprocess.run(
        [sys.executable, str(_pptx_generate_script()), "--output", str(out_path)],
        input=json.dumps(spec), capture_output=True, text=True, timeout=30,
    )
    assert proc.returncode == 0, proc.stderr
    prs = Presentation(str(out_path))
    assert len(prs.slides) == 2


def test_generate_pptx_script_degrades_gracefully_on_unknown_values(tmp_path):
    """Unknown layout/theme/icon values must never crash generation — same
    never-raises contract as pptx_themes.resolve_palette/resolve_icon."""
    import subprocess
    import sys

    from pptx import Presentation

    spec = {
        "title": "Weird Deck", "theme": "not_a_real_palette",
        "slides": [{"layout": "not_a_real_layout", "title": "Test", "bullets": ["a"], "icon": "not_a_real_icon"}],
    }
    out_path = tmp_path / "unknown.pptx"
    proc = subprocess.run(
        [sys.executable, str(_pptx_generate_script()), "--output", str(out_path)],
        input=json.dumps(spec), capture_output=True, text=True, timeout=30,
    )
    assert proc.returncode == 0, proc.stderr
    prs = Presentation(str(out_path))
    assert len(prs.slides) == 2
